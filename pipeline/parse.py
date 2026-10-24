"""Parse the district's monthly Excel lists into plain rows (stage 3 of the pipeline).

``parse_workbook(path)`` returns a ``ParsedFile`` dict (see docs/DATA_MODEL.md):
the detected kind (wz / cp / ws), the file title and printed period label, a layout id
and every data row with

* ``cells`` – canonical fields (``lp``, ``number``, ``date``, ``applicant``, ``investment``,
  ``address``, ``parcels``) holding the *raw* cell values, and
* ``raw`` – ``[header label as printed, value]`` pairs in column order (datetimes as ISO).

Design rules (the source is messy and changes layout over time):

* The header row is located by its "Lp"/"Lp." cell, never by a fixed row number.
* Columns are mapped to fields by header **text** (diacritics-insensitive keywords), never by
  position – WZ and CP layouts order their columns differently, and the 2025+ CP header
  "Streszczenie ustaleń decyzji\\nulica, dzielnica" is a mislabelled summary column.
* A second header row (merged parents such as "Nr decyzji i data jej wydania" → "nr"/"data")
  and the "1 2 3 4 5 6 7" column-numbering row are recognised and skipped.
* Data rows are scanned to the end of the sheet (gaps allowed).  Rows that contain nothing but
  a pre-printed Lp number are template padding and are reported in ``skipped`` (not warnings);
  so are table notes ("brak decyzji wydanych w grudniu") that stand in for missing data rows.
* Nothing with content is dropped silently: rows/columns that cannot be mapped are still
  emitted (with their raw values) together with a warning.
* The same rules read the city-wide lists of Biuro Architektury i Planowania Przestrzennego
  (BAiPP, layouts ``baipp-ws/-wz/-cp`` and the 2020 ``baipp-*-2020`` with 2-row headers);
  filtering them to Praga-Południe is not the parser's job (see :mod:`pipeline.baipp`).

Downloaded workbooks are untrusted input: openpyxl only reads cell values (``data_only``),
nothing from the file is ever executed.
"""

from __future__ import annotations

import re
import warnings as _warnings
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from .textutil import cell_to_json, collapse_ws, fold, is_blank, month_from_text

__all__ = ["FIELDS", "classify_header", "cp_series", "detect_kind", "is_baipp_title", "parse_workbook"]

#: Canonical field names produced in ``row["cells"]``.
FIELDS = ("lp", "number", "date", "applicant", "investment", "address", "parcels")

_LP_RE = re.compile(r"^\s*l\.?\s*p\.?\s*$", re.I)
_LABEL_SEP = " › "  # parent › child label for 2-row headers
_MAX_HEADER_SCAN = 40  # header must be within the first N rows of a sheet

# Header-text rules: (field, regex over folded+collapsed label).  Order matters: the first
# matching rule wins, so e.g. "streszczenie ustalen decyzji ulica, dzielnica" (2025+ CP layout,
# mislabelled) maps to the summary, not to the address.
_HEADER_RULES: list[tuple[str, re.Pattern[str]]] = [
    ("lp", re.compile(r"^l\.?\s*p\.?$")),
    ("investment", re.compile(r"streszczeni|nazwa inwestycji|rodzaj inwestycji|opis inwestycji")),
    ("applicant", re.compile(r"wnioskodawc|inwestor")),
    ("number", re.compile(r"znak sprawy|numer decyzji|^nr decyzji$|^nr\.?$|^numer$|sygnatur")),
    ("date", re.compile(r"^data\b|data wydania|data wszczecia|data decyzji")),
    ("address", re.compile(r"adres|^ulica|ulica, dzielnica")),
    ("parcels", re.compile(r"dzialk|\bdz\.|obreb|oznaczenie nieruchomosci|nieruchomosc")),
]

# Sub-header cells look like these (old 2020–2024 layout, row 5).
_SUBHEADER_RE = re.compile(r"^(nr\b|data\b|ulica|dzielnica|nr dz|dz\.|obreb)")


def classify_header(label: str) -> str | None:
    """Map a header label (any case/diacritics/whitespace) to a canonical field name.

    For 2-row headers pass the most specific text first (the child), e.g. ``"nr"``.

    >>> classify_header("Nazwa  wnioskodawcy")
    'applicant'
    >>> classify_header("Streszczenie ustaleń decyzji\\nulica, dzielnica")
    'investment'
    >>> classify_header("Oznaczenie nieruchomości\\nnr dz. ew., obręb")
    'parcels'
    """
    text = collapse_ws(fold(label))
    if not text:
        return None
    for field, rx in _HEADER_RULES:
        if rx.search(text):
            return field
    return None


def detect_kind(*texts: str) -> str | None:
    """Detect the list kind from title / sheet-name / header texts.

    ``ws`` (initiated proceedings) is checked first because its title also mentions both
    "warunkach zabudowy" and "celu publicznego".
    """
    folded = collapse_ws(fold(" ".join(t for t in texts if t)))
    if not folded:
        return None
    if re.search(r"wszcz[eę]t|wszczeci|znak sprawy", folded):
        return "ws"
    if re.search(r"cel\w* publiczn|lokalizacji inwestycji|ulicp", folded):
        return "cp"
    if re.search(r"warunk\w* zabudow|\bwz\b", folded):
        return "wz"
    return None


# --------------------------------------------------------------------------- helpers


def _norm_label(value: Any) -> str:
    return collapse_ws(cell_to_json(value)) if not is_blank(value) else ""


def _is_numbering_row(values: Sequence[Any]) -> bool:
    """The "1 2 3 4 5 6 7" row under the header (sometimes with a gap, e.g. 1 2 _ 4 5 6 7, or without
    its first number, e.g. _ 2 3 4 5 6 7 in the BAiPP WZ lists)."""
    nums = []
    for v in values:
        if is_blank(v):
            continue
        if isinstance(v, bool):
            return False
        if isinstance(v, (int, float)) and float(v).is_integer():
            nums.append(int(v))
        elif isinstance(v, str) and v.strip().isdigit():
            nums.append(int(v.strip()))
        else:
            return False
    return (len(nums) >= 4 and nums == sorted(nums) and len(set(nums)) == len(nums) and nums[0] in (0, 1, 2)
            and nums[-1] <= 30)


# A table note in place of data rows: "brak decyzji wydanych w grudniu", "w sierpniu i wrześniu 2026 r. nie
# wydano decyzji" (BAiPP lists). Only a row with a single text cell and no Lp/number/date qualifies.
_NOTE_RE = re.compile(r"\bbrak\b|\bnie\s+(?:wydano|wszczeto|bylo|zlozono)\b")


def _note_text(values: Sequence[Any], layout: _SheetLayout) -> str | None:
    """The text of a table-note row (see ``_NOTE_RE``), else ``None``."""
    filled = [(j, v) for j, v in enumerate(values) if not is_blank(v)]
    if len(filled) != 1 or not isinstance(filled[0][1], str):
        return None
    j, text = filled[0]
    if layout.fields.get(j) in ("lp", "number", "date"):
        return None
    text = collapse_ws(text)
    return text if len(text) <= 200 and _NOTE_RE.search(fold(text)) else None


def _merged_parent_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """(row, col) → top-left (row, col) for every cell covered by a merged range (1-based)."""
    out: dict[tuple[int, int], tuple[int, int]] = {}
    merged = getattr(ws, "merged_cells", None)  # absent in read-only mode
    for rng in (merged.ranges if merged is not None else ()):
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = (rng.min_row, rng.min_col)
    return out


def _find_header(rows: list[tuple[Any, ...]]) -> int | None:
    """Index (0-based) of the first row having an ``Lp`` cell and at least two more text cells."""
    for i, row in enumerate(rows[:_MAX_HEADER_SCAN]):
        if any(isinstance(v, str) and _LP_RE.match(v) for v in row):
            texts = [v for v in row if isinstance(v, str) and v.strip()]
            if len(texts) >= 3:
                return i
    return None


def _period_label(texts: Iterable[str]) -> str | None:
    """The reporting month as printed (e.g. "Lipiec 2026", "SIERPIEŃ", "LUTY-21", "maj 2023r.")."""
    for text in texts:
        line = collapse_ws(text)
        if month_from_text(line) is None:
            continue
        folded = fold(line)
        # Old layout: "Wydział … dla Dzielnicy Praga-Południe SIERPIEŃ……" → keep the tail.
        m = re.search(r"praga\s*-?\s*poludni\w*", folded)
        tail = line[m.end():] if m else line
        tail = re.sub(r"[…\.]{2,}.*$", "", tail)          # dotted fill-in lines
        tail = re.sub(r"/\s*nazwa organu\s*/", "", tail, flags=re.I)
        tail = tail.strip(" -–—…")
        if month_from_text(tail) is not None and len(tail) <= 40:
            return tail
        if len(line) <= 40:
            return line
    return None


def _title(texts: list[str]) -> str | None:
    """The list's main title ("Informacja o wydanych decyzjach …")."""
    for text in texts:
        if re.search(r"informacj|rejestr", fold(text)):
            return collapse_ws(text)
    for text in texts:
        if not re.search(r"zalacznik|wydzial|zarzad|urzad", fold(text)) and month_from_text(text) is None:
            return collapse_ws(text)
    return None


@dataclass
class _SheetLayout:
    """Column structure of one table (header [+ sub-header] [+ numbering row])."""

    labels: dict[int, str]      # col index (0-based) → printed label ("parent › child" for 2-row headers)
    fields: dict[int, str]      # col index → canonical field
    data_start: int             # 0-based row index of the first data row
    has_subheader: bool
    warnings: list[str]


def _build_layout(ws, rows: list[tuple[Any, ...]], h: int, sheet: str) -> _SheetLayout:
    """Work out labels/fields for the header found at row index ``h``."""
    warnings: list[str] = []
    header = list(rows[h])
    width = max(len(r) for r in rows) if rows else len(header)
    header += [None] * (width - len(header))
    merged = _merged_parent_map(ws)

    sub: list[Any] | None = None
    nxt = h + 1
    if nxt < len(rows):
        cand = list(rows[nxt]) + [None] * (width - len(rows[nxt]))
        lp_col = next((j for j, v in enumerate(header) if isinstance(v, str) and _LP_RE.match(v)), None)
        texts = [(j, v) for j, v in enumerate(cand) if isinstance(v, str) and v.strip()]
        lp_empty = lp_col is None or is_blank(cand[lp_col])
        if texts and lp_empty and all(_SUBHEADER_RE.match(collapse_ws(fold(v))) for _, v in texts):
            sub = cand
            nxt += 1

    labels: dict[int, str] = {}
    specific: dict[int, str] = {}   # text used for field classification
    for j in range(width):
        parent = _norm_label(header[j])
        if not parent:
            # covered by a merged parent? (e.g. "Nr decyzji i data jej wydania" over C4:D4)
            top = merged.get((h + 1, j + 1))
            if top and top != (h + 1, j + 1):
                parent = _norm_label(rows[top[0] - 1][top[1] - 1]) if top[0] - 1 < len(rows) else ""
            elif sub is not None and not is_blank(sub[j]):
                # no merge info: inherit from the nearest labelled column to the left
                k = j - 1
                while k >= 0 and not _norm_label(header[k]):
                    k -= 1
                parent = _norm_label(header[k]) if k >= 0 else ""
        child = _norm_label(sub[j]) if sub is not None else ""
        if parent and child:
            labels[j] = f"{parent}{_LABEL_SEP}{child}"
            specific[j] = child
        elif parent or child:
            labels[j] = parent or child
            specific[j] = parent or child

    fields: dict[int, str] = {}
    taken: dict[str, int] = {}
    for j in sorted(specific):
        field = classify_header(specific[j])
        if field is None and specific[j] != labels[j]:
            field = classify_header(labels[j])
        if field is None:
            warnings.append(f"arkusz „{sheet}”: kolumna {get_column_letter(j + 1)} "
                            f"(„{labels[j]}”) nie została rozpoznana – zachowano ją tylko w danych źródłowych")
            continue
        if field in taken:
            warnings.append(f"arkusz „{sheet}”: kolumny {get_column_letter(taken[field] + 1)} i "
                            f"{get_column_letter(j + 1)} wyglądają na to samo pole ({field}); użyto pierwszej")
            continue
        taken[field] = j
        fields[j] = field

    missing = [f for f in FIELDS if f not in taken]
    if missing:
        warnings.append(f"arkusz „{sheet}”: brak kolumn: {', '.join(missing)}")

    if nxt < len(rows) and _is_numbering_row(rows[nxt]):
        nxt += 1
    return _SheetLayout(labels, fields, nxt, sub is not None, warnings)


def is_baipp_title(title_texts: Iterable[str]) -> bool:
    """True for lists of the city's Biuro Architektury i Planowania Przestrzennego (BAiPP)."""
    return any(re.search(r"biur\w*\s+architektury", collapse_ws(fold(t))) for t in title_texts)


def _layout_id(kind: str | None, title_texts: list[str], has_subheader: bool, lp_col: int | None) -> str:
    annex = any("zalacznik" in fold(t) for t in title_texts)  # 2020 "Załącznik nr N do zarządzenia 1129/2020"
    if is_baipp_title(title_texts):
        return f"baipp-{kind or 'unknown'}" + ("-2020" if has_subheader or annex else "")
    old = has_subheader or annex or (lp_col or 0) > 0
    return f"{'old' if old else 'new'}-{kind or 'unknown'}"


def cp_series(*texts: str) -> str | None:
    """Series of a ULICP list from its title: ``c1`` (krajowe/wojewódzkie) or ``c2`` (powiatowe/gminne)."""
    folded = collapse_ws(fold(" ".join(t for t in texts if t)))
    local = bool(re.search(r"powiatow|gminn", folded))
    national = bool(re.search(r"krajow|wojewodzk", folded))
    if local == national:
        return None
    return "c2" if local else "c1"


# --------------------------------------------------------------------------- main entry


def parse_workbook(path: str | Path) -> dict[str, Any]:
    """Parse one workbook into a ``ParsedFile`` dict (see module docstring / DATA_MODEL.md).

    Never raises on odd content; problems are reported in ``warnings``.  Unreadable files
    (not an xlsx at all) do raise – the caller decides how to report those.
    """
    path = Path(path)
    with _warnings.catch_warnings():
        _warnings.simplefilter("ignore")  # openpyxl chatters about styles/extensions
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)

    result: dict[str, Any] = {
        "kind": None,
        "title": None,
        "period_label": None,
        "layout": None,
        "series": None,
        "sheets": [],
        "rows": [],
        "skipped": [],
        "warnings": [],
    }
    kinds_seen: list[str] = []

    try:
        for ws in wb.worksheets:
            sheet = ws.title
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            nonempty = [(i, r) for i, r in enumerate(rows) if any(not is_blank(v) for v in r)]
            sheet_info: dict[str, Any] = {"name": sheet, "max_row": ws.max_row, "header_row": None,
                                          "kind": None, "layout": None, "rows": 0}
            result["sheets"].append(sheet_info)
            if not nonempty:
                continue

            h = _find_header(rows)
            if h is None:
                # Content without a recognisable table: keep it, flag it.
                result["warnings"].append(
                    f"arkusz „{sheet}”: nie znaleziono wiersza nagłówka (Lp) – "
                    f"{len(nonempty)} niepustych wierszy zachowano bez mapowania pól")
                for i, r in nonempty:
                    raw = [[f"(kolumna {get_column_letter(j + 1)})", cell_to_json(v)]
                           for j, v in enumerate(r) if not is_blank(v)]
                    result["rows"].append({"sheet": sheet, "row": i + 1, "cells": {f: None for f in FIELDS},
                                           "raw": raw, "warnings": ["wiersz spoza rozpoznanej tabeli"]})
                continue

            title_texts = [collapse_ws(v) for r in rows[:h] for v in r if isinstance(v, str) and v.strip()]
            header_texts = [collapse_ws(v) for v in rows[h] if isinstance(v, str)]
            kind = detect_kind(*title_texts) or detect_kind(sheet) or detect_kind(*header_texts)
            header_kind = "ws" if detect_kind(*header_texts) == "ws" else None
            if kind and header_kind and kind != header_kind:
                result["warnings"].append(
                    f"arkusz „{sheet}”: tytuł wskazuje „{kind}”, a nagłówki kolumn „{header_kind}” – przyjęto nagłówki")
                kind = header_kind
            if kind is None:
                result["warnings"].append(f"arkusz „{sheet}”: nie rozpoznano rodzaju listy (WZ/CP/wszczęte)")

            layout = _build_layout(ws, rows, h, sheet)
            result["warnings"].extend(layout.warnings)
            lp_col = next((j for j, f in layout.fields.items() if f == "lp"), None)
            sheet_info.update(header_row=h + 1, kind=kind,
                              layout=_layout_id(kind, title_texts, layout.has_subheader, lp_col))
            if kind:
                kinds_seen.append(kind)
            if result["title"] is None:
                result["title"] = _title(title_texts)
            if result["period_label"] is None:
                result["period_label"] = _period_label(title_texts)
            if result["layout"] is None:
                result["layout"] = sheet_info["layout"]
            if kind == "cp" and result["series"] is None:
                result["series"] = cp_series(*title_texts)

            current = layout
            for i in range(layout.data_start, len(rows)):
                r = rows[i]
                if all(is_blank(v) for v in r):
                    continue
                # A repeated header (second table on the same sheet) re-maps the columns.
                if _find_header([r]) == 0:
                    current = _build_layout(ws, rows, i, sheet)
                    result["warnings"].extend(current.warnings)
                    result["warnings"].append(f"arkusz „{sheet}”: kolejny nagłówek tabeli w wierszu {i + 1}")
                    continue
                if i < current.data_start:
                    continue  # sub-header / numbering row of a repeated header
                note = _note_text(r, current)
                if note is not None:
                    result["skipped"].append({"sheet": sheet, "row": i + 1, "reason": f"adnotacja: „{note}”"})
                    continue
                row = _make_row(sheet, i, r, current)
                if row is None:
                    result["skipped"].append({"sheet": sheet, "row": i + 1, "reason": "tylko numer Lp"})
                    continue
                result["rows"].append(row)
                sheet_info["rows"] += 1
    finally:
        wb.close()

    distinct = sorted(set(kinds_seen))
    if len(distinct) > 1:
        result["warnings"].append(f"arkusze w pliku mają różne rodzaje list: {', '.join(distinct)}")
    result["kind"] = kinds_seen[0] if kinds_seen else None
    return result


def _make_row(sheet: str, i: int, values: tuple[Any, ...], layout: _SheetLayout) -> dict[str, Any] | None:
    """Build a ParsedFile row, or None when the row only carries a pre-printed Lp number."""
    cells: dict[str, Any] = {f: None for f in FIELDS}
    raw: list[list[Any]] = []
    row_warnings: list[str] = []
    content = False
    width = max(len(values), max(layout.labels, default=-1) + 1)
    for j in range(width):
        v = values[j] if j < len(values) else None
        blank = is_blank(v)
        field = layout.fields.get(j)
        label = layout.labels.get(j)
        if field is not None:
            cells[field] = None if blank else v
            if field != "lp" and not blank:
                content = True
        elif not blank:
            content = True
            if label is None:
                label = f"(kolumna {get_column_letter(j + 1)})"
            row_warnings.append(f"wartość w kolumnie bez rozpoznanego pola: „{label}”")
        if label is not None:
            raw.append([label, None if blank else cell_to_json(v)])
    if not content:
        return None
    return {"sheet": sheet, "row": i + 1, "cells": cells, "raw": raw, "warnings": row_warnings}
