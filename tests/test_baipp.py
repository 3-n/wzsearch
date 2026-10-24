"""Tests for the BAiPP source (city architecture bureau lists): crawler, parser, district filter,
normaliser specifics, out-of-district parcels and the build integration.

Fixtures in tests/fixtures/baipp/ are real files from architektura.um.warszawa.pl (fetched 2026-10-23):

* html/ – the current-year page and two yearly archive pages, <script>/<style>/<svg> stripped;
* xlsx/ – seven of the 28 published workbooks, one per layout / edge case (the file name says which).
"""

from __future__ import annotations

import io
import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
import requests

from pipeline import baipp, build
from pipeline import crawl as crawl_mod
from pipeline import geocode as gc
from pipeline.crawl import (
    BAIPP_URL,
    baipp_series,
    crawl_baipp,
    detect_baipp_category,
    parse_baipp_page,
    run_crawl,
)
from pipeline.normalize import merge_duplicates, normalize_row, parse_address, parse_parcels
from pipeline.parse import parse_workbook

ROOT = Path(__file__).resolve().parents[1]

FIX = ROOT / "tests" / "fixtures" / "baipp"
HTML = FIX / "html"
XLSX = FIX / "xlsx"
SITE = "https://architektura.um.warszawa.pl"
XLSX_MAGIC = b"PK\x03\x04"


def stub_xlsx(file_id: str) -> bytes:
    """A small valid zip archive, different for every file (downloads are checked whole)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", f"<Types id='{file_id}'/>")
    return buf.getvalue()
NOW = "2026-10-23T00:00:00Z"


def html(name: str) -> str:
    return (HTML / f"{name}.html").read_text(encoding="utf-8")


def fx(name: str) -> Path:
    return XLSX / f"{name}.xlsx"


def current_page():
    return parse_baipp_page(html("postepowania-i-decyzje"), BAIPP_URL, today="2026-10-23")


# =============================================================================
# Crawler: pages, blocks, archive discovery
# =============================================================================


def test_current_page_blocks_categories_series_and_period():
    page, archive = current_page()
    assert (page["page_id"], page["source"], page["period"], page["published_at"]) == (
        "postepowania-i-decyzje", "baipp", "2026", None)
    assert page["title"] == "Informacja o wszczętych postępowaniach i wydanych decyzjach administracyjnych"
    assert [(b["category"], b["series"], b["declared_count"]) for b in page["blocks"]] == [
        ("ws", None, None), ("wz", None, None), ("cp", "c2", None), ("cp", "c1", None)]
    assert all(len(b["file_ids"]) == 1 for b in page["blocks"])
    assert page["blocks"][1]["text"] == "2. Informacja o wydanych decyzjach w sprawie warunków zabudowy IX.2026 (xls, 31 kB)"
    assert page["warnings"] == []
    link = page["links"][1]
    assert link["url"] == (SITE + "/documents/12025039/171613001/informacja_o_wydanych_decyzjach_WZ_IX_2026+%2B.xlsx/"
                                  "1d75d6d6-239a-2dc1-843b-70184b23709c?t=1791798490945")
    assert link["file_id"] == crawl_mod.file_id_of(link["document_path"])  # file_id rule unchanged
    assert link["filename"] == "informacja_o_wydanych_decyzjach_WZ_IX_2026 +.xlsx"
    # the "Dane z ubiegłych lat" box: discovered, irregular slugs (never guessed)
    assert [s["page_id"] for s in archive] == ["rejestry2020", "rejestry2021", "rejestry2022", "rejestry23",
                                              "rejestry24", "rejestry25"]
    assert archive[3] == {"url": SITE + "/-/rejestry23", "page_id": "rejestry23",
                          "title": "Informacja o wszczętych postępowaniach i wydanych decyzjach w 2023 r."}


def test_yearly_pages_period_from_title_and_publication_time():
    page, archive = parse_baipp_page(html("rejestry2020"), SITE + "/-/rejestry2020")
    assert (page["page_id"], page["period"], page["published_at"]) == ("rejestry2020", "2020", "2022-02-16T14:09:00Z")
    assert page["title"] == "Informacja o wszczętych postępowaniach i wydanych decyzjach (sierpień - grudzień 2020 r.)"
    assert [b["category"] for b in page["blocks"]] == ["ws", "wz", "cp", "cp"]
    assert [b["series"] for b in page["blocks"]] == [None, None, "c2", "c1"]
    assert archive is None and page["warnings"] == []
    page, _ = parse_baipp_page(html("rejestry25"), SITE + "/-/rejestry25")
    assert (page["period"], page["published_at"]) == ("2025", "2026-01-14T11:40:00Z")


@pytest.mark.parametrize("text, category, series", [
    ("1. Informacja o wszczętych postępowaniach o warunkach zabudowy i ustaleniu lokalizacji inwestycji celu "
     "publicznego IX.2026 (xls, 37 kB)", "ws", None),
    ("2. Informacja o wydanych decyzjach w sprawie warunków zabudowy - I-XII 2025 (xml, 35KB)", "wz", None),
    ("3. Informacja o wydanych decyzjach w spr. ust. lokalizacji inwestycji celu publ. o znaczeniu powiatowym i "
     "gminnym- VIII-XII 2020", "cp", "c2"),
    ("4. Informacja o wydanych decyzjach w spr.ust. lokalizacji inwestycji celu publ. o znaczeniu krajowym i "
     "wojewódzkim-VIII-XII.2020", "cp", "c1"),
    ("informacja_dec._licp_pow_gmi_sierpien-grudzien_20.xlsx", "cp", "c2"),
    ("informacja o wydanych decyzjach ws. ustalenia_licp_kraj-woj_XII_23+.xlsx", "cp", "c1"),
    ("Zarządzenie nr 1129/2020", None, None),
])
def test_attachment_category_and_series(text, category, series):
    assert detect_baipp_category(text) == category
    if category == "cp":
        assert baipp_series(text) == series


def test_page_without_year_or_attachments_is_flagged():
    page, archive = parse_baipp_page('<div class="journal-content-article" data-analytics-asset-title="Rejestry">'
                                     "<p>brak</p></div>", SITE + "/-/rejestry", today="2027-01-05")
    assert page["period"] == "2027" and page["blocks"] == [] and archive is None
    assert any("nie znaleziono załączników" in w for w in page["warnings"])
    assert any("przyjęto bieżący rok (2027)" in w for w in page["warnings"])


# --------------------------------------------------------------------------- fake network


class FakeResponse:
    def __init__(self, url, status=200, content=b"", headers=None):
        self.url, self.status_code = url, status
        self.content = content.encode("utf-8") if isinstance(content, str) else content
        self.headers = headers or {}
        self.encoding = None

    @property
    def text(self) -> str:
        return self.content.decode(self.encoding or "utf-8")

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code} for {self.url}", response=self)


class FakeSession:
    """Canned responses by exact URL; unknown URLs → 404."""

    def __init__(self, routes=None):
        self.routes = dict(routes or {})
        self.calls: list[str] = []

    def get(self, url, **kwargs):
        self.calls.append(url)
        route = self.routes.get(url)
        if isinstance(route, Exception):
            raise route
        if isinstance(route, FakeResponse):
            return route
        if route is None:
            return FakeResponse(url, 404, b"not found")
        if isinstance(route, bytes):
            return FakeResponse(url, 200, route, {"Last-Modified": "Mon, 12 Oct 2026 09:48:11 GMT"})
        return FakeResponse(url, 200, route, {"Content-Type": "text/html;charset=UTF-8"})


def yearly_page(year: int, files: list[tuple[str, str]]) -> str:
    """A minimal archive page with the real markup skeleton (article + attachment list)."""
    items = "".join(f'<li><a href="{href}" class="attachment-link"><span class="attachment-name">{text}</span></a></li>'
                    for href, text in files)
    return (f'<div class="journal-content-article" data-analytics-asset-title="Informacja o wszczętych '
            f'postępowaniach i wydanych decyzjach ({year} r.)"><article><time datetime="{year + 1}-01-14T11:40:00Z">'
            f'x</time><div class="article-description"><p>Publikujemy</p></div><div class="article-attachments">'
            f'<ul>{items}</ul></div></article></div>')


def baipp_session(main_html: str | None = None) -> FakeSession:
    """The current page + 2025 + 2020 from fixtures, 2021–2024 as minimal synthetic pages, every file a stub."""
    session = FakeSession({BAIPP_URL: main_html or html("postepowania-i-decyzje"),
                           SITE + "/-/rejestry25": html("rejestry25"), SITE + "/-/rejestry2020": html("rejestry2020")})
    for slug, year in (("rejestry2021", 2021), ("rejestry2022", 2022), ("rejestry23", 2023), ("rejestry24", 2024)):
        session.routes[f"{SITE}/-/{slug}"] = yearly_page(year, [
            (f"/documents/1/{year}/wz_{year}.xlsx/0f0f0f0f-0000-0000-0000-00000000{year}?t=1", f"2. … zabudowy {year}")])
    for url in list(session.routes):
        if url.startswith(SITE + "/-/") or url == BAIPP_URL:
            page, _ = parse_baipp_page(session.routes[url], url)
            for link in page["links"]:
                session.routes[link["url"]] = stub_xlsx(link["file_id"])
    return session


def test_crawl_baipp_follows_the_archive_box():
    session = baipp_session()
    pages, complete = crawl_baipp(session)
    assert complete
    assert [p["page_id"] for p in pages] == ["postepowania-i-decyzje", "rejestry2020", "rejestry2021",
                                            "rejestry2022", "rejestry23", "rejestry24", "rejestry25"]
    assert [p["period"] for p in pages] == ["2026", "2020", "2021", "2022", "2023", "2024", "2025"]
    assert all(p["source"] == "baipp" for p in pages)
    assert SITE + "/-/rejestry2026" not in session.calls and SITE + "/-/rejestry26" not in session.calls


def test_crawl_baipp_failures_are_not_fatal_and_use_the_previous_manifest(tmp_path):
    out, raw = tmp_path / "manifest.json", tmp_path / "raw"
    first = run_crawl(out, raw, session=baipp_session(), sources=("baipp",))
    assert len(first["pages"]) == 7 and len(first["files"]) == 4 + 4 + 4 * 1 + 4

    session = baipp_session()
    session.routes[SITE + "/-/rejestry25"] = FakeResponse(SITE + "/-/rejestry25", 503)
    manifest = run_crawl(out, raw, session=session, sources=("baipp",))
    pages = {p["page_id"]: p for p in manifest["pages"]}
    assert "poprzedniej wersji" in " ".join(pages["rejestry25"]["warnings"])
    assert len(pages["rejestry25"]["blocks"]) == 4 and len(manifest["files"]) == len(first["files"])

    down = FakeSession({BAIPP_URL: requests.ConnectionError("down")})
    manifest = run_crawl(out, raw, session=down, sources=("baipp",))
    assert {p["page_id"] for p in manifest["pages"]} == {p["page_id"] for p in first["pages"]}
    assert not any(p.get("removed_from_source_at") for p in manifest["pages"])  # incomplete crawl: nothing removed
    assert not any(f.get("removed_from_source_at") for f in manifest["files"])
    assert "poprzedniej wersji" in " ".join(next(p for p in manifest["pages"]
                                                 if p["page_id"] == "postepowania-i-decyzje")["warnings"])


def test_run_crawl_keeps_every_monthly_snapshot_of_the_current_year(tmp_path):
    out, raw = tmp_path / "manifest.json", tmp_path / "raw"
    september = run_crawl(out, raw, session=baipp_session(), sources=("baipp",))
    files = {f["file_id"]: f for f in september["files"]}
    sept_wz = next(f for f in september["files"] if f["filename"] == "informacja_o_wydanych_decyzjach_WZ_IX_2026 +.xlsx")
    assert (sept_wz["source"], sept_wz["category"], sept_wz["period"], sept_wz["page_ids"]) == (
        "baipp", "wz", "2026", ["postepowania-i-decyzje"])
    cp_series = {f["series"] for f in files.values() if f["category"] == "cp" and f["period"] == "2026"}
    assert cp_series == {"c1", "c2"}

    # October: the page links a new WZ snapshot (new folder + uuid), the September link disappears
    october_html = html("postepowania-i-decyzje").replace(
        "/documents/12025039/171613001/informacja_o_wydanych_decyzjach_WZ_IX_2026+%2B.xlsx/"
        "1d75d6d6-239a-2dc1-843b-70184b23709c?t=1791798490945",
        "/documents/12025039/180000000/informacja_o_wydanych_decyzjach_WZ_X_2026+%2B.xlsx/"
        "2e86e7e7-0000-0000-0000-000000000001?t=1794400000000").replace("IX.2026 (xls, 31 kB)", "X.2026 (xls, 32 kB)")
    session = baipp_session(october_html)
    october = run_crawl(out, raw, session=session, sources=("baipp",))
    by_id = {f["file_id"]: f for f in october["files"]}
    assert len(october["files"]) == len(september["files"]) + 1
    kept = by_id[sept_wz["file_id"]]
    assert kept["removed_from_source_at"] and kept["sha256"] == sept_wz["sha256"]
    assert kept["page_ids"] == ["postepowania-i-decyzje"] and (raw / f"{kept['file_id']}.xlsx").is_file()
    assert sept_wz["url"] not in session.calls  # kept copy, nothing downloaded again
    new = next(f for f in october["files"] if "X_2026" in f["filename"] and "IX" not in f["filename"])
    assert not new.get("removed_from_source_at") and new["period"] == "2026"
    linked = {fid for p in october["pages"] for b in p["blocks"] for fid in b["file_ids"]}
    assert sept_wz["file_id"] not in linked and new["file_id"] in linked

    # a snapshot linked again loses the mark; a third run keeps the first "removed" date
    again = run_crawl(out, raw, session=baipp_session(october_html), sources=("baipp",))
    assert {f["file_id"]: f.get("removed_from_source_at") for f in again["files"]}[sept_wz["file_id"]] == \
        kept["removed_from_source_at"]
    back = run_crawl(out, raw, session=baipp_session(), sources=("baipp",))
    assert not {f["file_id"]: f for f in back["files"]}[sept_wz["file_id"]].get("removed_from_source_at")


def test_run_crawl_skip_flag_carries_the_other_source_over(tmp_path):
    out, raw = tmp_path / "manifest.json", tmp_path / "raw"
    first = run_crawl(out, raw, session=baipp_session(), sources=("baipp",))
    session = FakeSession()
    with pytest.raises(requests.HTTPError):  # the district index is unreachable here
        run_crawl(out, raw, session=session)
    again = run_crawl(out, raw, session=baipp_session(), sources=("baipp",))
    assert [p["page_id"] for p in again["pages"]] == [p["page_id"] for p in first["pages"]]
    with pytest.raises(ValueError):
        run_crawl(out, raw, session=session, sources=("rwdz",))


def test_cli_skip_flags(tmp_path, monkeypatch):
    monkeypatch.setattr(crawl_mod, "make_session", lambda **kw: baipp_session())
    code = crawl_mod.main(["--out", str(tmp_path / "m.json"), "--raw", str(tmp_path / "raw"), "--skip-district", "-q"])
    assert code == 0
    manifest = json.loads((tmp_path / "m.json").read_text(encoding="utf-8"))
    assert {p["source"] for p in manifest["pages"]} == {"baipp"} and len(manifest["files"]) == 16


# =============================================================================
# Parser: every BAiPP layout, no dropped rows
# =============================================================================


def _oracle_rows(path: Path) -> int:
    """Independent count: rows below the header whose number column holds a decision number or case sign."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if any(isinstance(v, str) and v.strip().rstrip(".").lower() in
                                                       ("lp", "l.p") for v in r))
    return sum(1 for r in rows[header + 1:] if isinstance(r[2], str)
               and re.search(r"\d+\s*/|\.\d{4}\.\d+\.\d{4}", r[2]))


BAIPP_FIXTURES = [
    # name, kind, layout, series, rows, skipped reasons, Praga-Południe rows
    ("baipp-ws-2026-09-ytd-ezd-signs", "ws", "baipp-ws", None, 85, [], 2),
    ("baipp-wz-2026-09-ytd-numbering-without-1", "wz", "baipp-wz", None, 101, [], 1),
    ("baipp-cp-c2-2026-09-note-row", "cp", "baipp-cp", "c2", 6,
     ["adnotacja: „w sierpniu i wrześniu 2026 r. nie wydano decyzji”"], 0),
    ("baipp-wz-2020-merged-header", "wz", "baipp-wz-2020", None, 79, [], 4),
    ("baipp-cp-c1-2020-13-columns", "cp", "baipp-cp-2020", "c1", 6, [], 0),
    ("baipp-cp-c1-2021-note-row", "cp", "baipp-cp", "c1", 5, ["adnotacja: „brak decyzji wydanych w grudniu”"], 1),
    ("baipp-cp-c2-2022-cut-obreb", "cp", "baipp-cp", "c2", 13, [], 6),
]


@pytest.mark.parametrize("name, kind, layout, series, n_rows, skipped, n_pp", BAIPP_FIXTURES)
def test_baipp_workbooks(name, kind, layout, series, n_rows, skipped, n_pp):
    pf = parse_workbook(fx(name))
    assert (pf["kind"], pf["layout"], pf["series"]) == (kind, layout, series)
    assert len(pf["rows"]) == n_rows == _oracle_rows(fx(name))
    assert [s["reason"] for s in pf["skipped"]] == skipped
    assert pf["warnings"] == [] and all(not r["warnings"] for r in pf["rows"])
    assert all(r["cells"]["number"] and r["cells"]["date"] for r in pf["rows"])
    assert sum(baipp.in_district(r["cells"]["address"]) for r in pf["rows"]) == n_pp


def test_two_row_header_and_numbering_row_without_its_first_number():
    pf = parse_workbook(fx("baipp-wz-2020-merged-header"))
    assert pf["rows"][0]["row"] == 7  # header 4–5, numbering row 6
    assert [label for label, _ in pf["rows"][0]["raw"]][:3] == [
        "Lp.", "Nr decyzji i data jej wydania › nr", "Nr decyzji i data jej wydania › data (rrrr-mm-dd)"]
    pf = parse_workbook(fx("baipp-wz-2026-09-ytd-numbering-without-1"))  # row 5 is "_ 2 3 4 5 6 7"
    assert pf["rows"][0]["row"] == 6 and pf["rows"][0]["cells"]["number"].startswith("1/AM/URN/WZ/2026/P")
    assert [label for label, _ in pf["rows"][0]["raw"]] == [
        "L.P.", "Nr decyzji", "data (rrrr-mm-dd)", "Nazwa wnioskodawcy", "Oznaczenie nieruchomości (ulica, dzielnica)",
        "Oznaczenie nieruchomości (nr działki ewid., obręb)", "Streszczenie ustaleń decyzji"]


def _raw_baipp_files() -> list[tuple[Path, dict]]:
    manifest = ROOT / "data" / "manifest.json"
    if not manifest.is_file():
        return []
    files = json.loads(manifest.read_text(encoding="utf-8")).get("files", [])
    return [(ROOT / f["local_path"], f) for f in files
            if f.get("source") == "baipp" and f.get("local_path") and (ROOT / f["local_path"]).is_file()]


@pytest.mark.skipif(not _raw_baipp_files(), reason="no BAiPP files downloaded yet")
def test_every_downloaded_baipp_workbook_parses_without_dropping_rows():
    for path, entry in _raw_baipp_files():
        pf = parse_workbook(path)
        assert pf["layout"].startswith("baipp-") and pf["kind"] == entry["category"], path.name
        assert len(pf["rows"]) == _oracle_rows(path), path.name
        assert all(r["cells"]["number"] for r in pf["rows"]), path.name
        if pf["kind"] == "cp":
            assert pf["series"] == entry["series"], path.name


# =============================================================================
# District filter
# =============================================================================


@pytest.mark.parametrize("address, keep", [
    ("Ostrobramska 73 \nPraga Południe \n\n", True),
    ("Chodakowska\nWawerska\nPraga Południe", True),
    ("Jana Nowaka Jeziorańskiego                Praga Południe", True),   # visual line break (padding)
    ("Władysława Umińskiego \nPraga- Południe", True),
    ("Jana Nowaka-Jeziorańskiego \nParga Południe", True),                 # typo in the 2021 list
    ("Grochowska 1\nPRAGA PŁD.", True),
    ("Wał Miedzeszyński 646 Praga-Poludnie", True),
    ("Sokola\nZamoście,\nWybrzeże Szczecińskie\nPraga Północ", False),
    ("Chełmżyńska\nRembertów", False),
    ("Grochowska\nPraga Południe\nWawer", False),                           # only the last line counts
    ("", False),
    (None, False),
])
def test_in_district_reads_the_last_line_of_the_address(address, keep):
    assert baipp.in_district(address) is keep


def test_applicant_naming_the_district_is_ignored():
    row = {"cells": {"applicant": "Spółdzielnia Spożywców Praga Południe", "address": "Wolska 1\nWola"}}
    assert not baipp.in_district(row["cells"]["address"])


def test_strip_district():
    assert baipp.strip_district("Chodakowska\nWawerska\nPraga Południe") == "Chodakowska\nWawerska"
    assert baipp.strip_district("Wał Miedzeszyński 646 Praga Południe") == "Wał Miedzeszyński 646"
    assert baipp.strip_district("Jana Nowaka-Jeziorańskiego \nParga Południe") == "Jana Nowaka-Jeziorańskiego"
    assert baipp.strip_district("Wolska 1\nWola") == "Wolska 1\nWola"


# =============================================================================
# Normaliser: numbers, outcomes, ids, applicants, addresses, parcels
# =============================================================================


@pytest.mark.parametrize("cell, kind, number, rest, others", [
    ("95/AM/PPD/WZ/2026/O\n\nodmowa ", "wz", "95/AM/PPD/WZ/2026/O", "odmowa", []),
    ("119/AM/PPD/WZ /2025/P\nprzeniesienie ", "wz", "119/AM/PPD/WZ/2025/P", "przeniesienie", []),
    ("99/AM/SRD/WZ/2026/W wygaszenie decyzji ", "wz", "99/AM/SRD/WZ/2026/W", "wygaszenie decyzji", []),
    ("10/AM/BEM/WZ/2026 /U\n", "wz", "10/AM/BEM/WZ/2026/U", "", []),
    ("89/AM/URN//WZ/2026/I\n\n", "wz", "89/AM/URN//WZ/2026/I", "", []),
    ("1/PRD/C2/22\numorzenie  licp", "cp", "1/PRD/C2/22", "umorzenie licp", []),
    ("wygaśnięcie decyzji + ", "wz", "wygaśnięciedecyzji+", "", []),          # no number at all: kept, id falls back
    ("AM-ZP-P.6730.14.2020\n", "ws", "AM-ZP-P.6730.14.2020", "", []),
    ("AM-ZP-P.6733.6.2022\nlicp", "ws", "AM-ZP-P.6733.6.2022", "licp", []),
    ("AM-ZP-N.6730.25.2026 \n\nEZD AM-ZP-N.6730.43.2026 ", "ws", "AM-ZP-N.6730.25.2026", "",
     ["AM-ZP-N.6730.43.2026"]),
    ("EZD AM-ZP-E.6733.1.2026", "ws", "AM-ZP-E.6733.1.2026", "", []),
])
def test_split_number(cell, kind, number, rest, others):
    parts = baipp.split_number(cell, kind)
    assert (parts.number, parts.rest, parts.others) == (number, rest, others)


@pytest.mark.parametrize("cell, kind, outcome, warning", [
    ("95/AM/PPD/WZ/2026/O\n\nodmowa ", "wz", "refusal", None),
    ("24/AM/PPD/WZ/2024/U\numorzenie ", "wz", "discontinued", None),
    ("33/AM/PPD/WZ/2025/P\nprzeniesienie", "wz", "transfer", None),
    ("3/AM/SRD/CP1/2026/W\n\n ", "cp", "expiry", None),
    ("101/AM/BEM/WZ/2026        \n", "wz", "granted", None),                  # 2024+: no suffix = granted
    ("67/AM/PPD/WZ/2024\nWZ", "wz", "granted", None),
    ("6/PRD/21\nwz", "wz", "granted", None),
    ("10/AM/PPD/CP2/2025\nlicp\n", "cp", "granted", None),
    ("1/PRD/C2/22\numorzenie  licp", "cp", "discontinued", None),
    ("147/PRD/23\nwygaśnięcie decyzji", "wz", "expiry", None),
    ("63/AM/URS/WZ/2024/P \n\nodmowa przeniesienia decyzji ", "wz", "refusal", None),  # refused transfer
    ("108/AM/WLO/WZ/2024/I\nodmowa zmiany ", "wz", "refusal", None),
    ("123/AM/WOL/WZ/2024/O\ninna \n", "wz", "refusal", "nierozpoznany opis rozstrzygnięcia („inna”)"),
    ("57/AM/BIE/WZ/2024/X\nuchylenie decyzji i zawieszenie postępowania ", "wz", None, "nie ustalono rozstrzygnięcia"),
    ("130/BIE/21\nstwierdzenie wydania decyzji z naruszeniem prawa ", "wz", None, "nie ustalono rozstrzygnięcia"),
    ("11/AM/SRD/WZ/2026 /I\n", "wz", None, "nieznane oznaczenie /I"),
    ("12/AM/X/WZ/2026/U\nprzeniesienie", "wz", "transfer", "wskazują różne rozstrzygnięcia"),
    ("90/ŻOL/20", "wz", None, None),                                            # old number, nothing stated
    ("AM-ZP-P.6730.14.2020", "ws", None, None),                                 # proceedings have no outcome
])
def test_parse_outcome(cell, kind, outcome, warning):
    parts = baipp.split_number(cell, kind)
    got, warnings = baipp.parse_outcome(parts.number, parts.rest, kind)
    assert got == outcome
    if warning:
        assert any(warning in w for w in warnings), warnings
    else:
        assert warnings == []


@pytest.mark.parametrize("value, expected", [
    ("dane osobowe ", "osoba fizyczna"),
    ("dane osobowe ⏎⏎".replace("⏎", "\n"), "osoba fizyczna"),
    ("Kancelaria Adwokacka \ndane osobowe", "Kancelaria Adwokacka (dane osobowe)"),
    ("Inicjatywa Projektowa \nsp. z o.o.\n\n", "Inicjatywa Projektowa sp. z o.o."),
])
def test_applicant_text(value, expected):
    assert baipp.applicant_text(value) == expected


@pytest.mark.parametrize("base, kind, number, series, expected", [
    ("wz-2026-95", "wz", "95/AM/PPD/WZ/2026/O", None, "wz-am-2026-95"),
    ("wz-2021-6", "wz", "6/PRD/21", None, "wz-am-2021-6"),
    ("cp-2022-1", "cp", "1/PRD/C2/22", "c2", "cp-am-2022-c2-1"),
    ("cp-2025-10", "cp", "10/AM/PPD/CP2/2025", None, "cp-am-2025-c2-10"),
    ("cp-2021-3", "cp", "3/PRD/21", "c1", "cp-am-2021-c1-3"),               # series from the file
    ("ws-6730-2020-14", "ws", "AM-ZP-P.6730.14.2020", None, "ws-am-6730-2020-14"),
    ("ws-6730-2025-5", "ws", "AM-ZP-E.6730.5.2025", None, "ws-am-zp-e-6730-2025-5"),  # per-unit numbering
    ("wz-f37a76419da-r12", "wz", "", None, "wz-am-f37a76419da-r12"),
])
def test_namespace_id(base, kind, number, series, expected):
    assert baipp.namespace_id(base, kind, number, series)[0] == expected


def test_namespace_id_series_conflict_is_reported():
    rid, warnings = baipp.namespace_id("cp-2023-10", "cp", "10/PRD/C2/23", "c1")
    assert rid == "cp-am-2023-c2-10" and any("seria decyzji" in w for w in warnings)


def baipp_records(name: str, year: str, file_id: str = "0123456789", series: str | None = None) -> list[dict]:
    pf = parse_workbook(fx(name))
    meta = {"file_id": file_id, "page_id": "postepowania-i-decyzje", "period": year, "dataset": "baipp",
            "series": series or pf["series"]}
    return [normalize_row(r, pf["kind"], meta) for r in pf["rows"] if baipp.in_district(r["cells"]["address"])]


def test_normalize_row_baipp_wz_refusal():
    (rec,) = baipp_records("baipp-wz-2026-09-ytd-numbering-without-1", "2026")
    assert rec["id"] == "wz-am-2026-95"
    assert (rec["dataset"], rec["issuer"], rec["outcome"], rec["filed_date"], rec["permit"]) == (
        "baipp", "city", "refusal", None, None)
    assert (rec["kind"], rec["procedure"], rec["number"]) == ("wz", "wz", "95/AM/PPD/WZ/2026/O")
    assert (rec["date"], rec["period"]) == ("2026-09-16", "2026-09")      # period = month of the decision
    assert (rec["applicant"], rec["applicant_type"]) == ("Inicjatywa Projektowa sp. z o.o.", "organization")
    assert rec["address"]["raw"] == "Ostrobramska 73 \nPraga Południe \n\n"
    assert (rec["address"]["display"], rec["address"]["street"], rec["address"]["number"]) == (
        "Ostrobramska 73", "Ostrobramska", "73")
    assert rec["parcels"] == [{"number": "7/5", "obreb": "3-05-24", "partial": True, "id": "146507_8.0524.7/5"}]
    src = rec["sources"][0]
    assert (src["file_id"], src["page_id"], src["row"]) == ("0123456789", "postepowania-i-decyzje", 103)
    assert src["raw"][1] == ["Nr decyzji", "95/AM/PPD/WZ/2026/O\nodmowa"]  # the source keeps the outcome word
    assert rec["warnings"] == []


def test_normalize_row_baipp_2020_person_and_parcel_groups():
    recs = {r["id"]: r for r in baipp_records("baipp-wz-2020-merged-header", "2020")}
    assert list(recs) == ["wz-am-2020-98", "wz-am-2020-110", "wz-am-2020-123", "wz-am-2020-154"]
    r98, r123 = recs["wz-am-2020-98"], recs["wz-am-2020-123"]
    assert (r98["applicant"], r98["applicant_type"], r98["outcome"]) == ("osoba fizyczna", "person", "discontinued")
    assert recs["wz-am-2020-110"]["number"] == "110/PRN/20"   # district code PRN: still Praga-Południe by address
    assert [(p["number"], p["obreb"]) for p in r123["parcels"]] == [
        ("3", "3-02-03"), ("2", "3-02-03"), ("8", "3-02-03"), ("9", "3-02-03"), ("10", "3-02-03"),
        ("11", "3-02-03"), ("12", "3-02-03"), ("9", "3-02-02"), ("12", "3-02-02"), ("21", "3-02-02")]
    assert r123["outcome"] == "granted" and r123["date"] == "2020-10-15"


def test_normalize_row_baipp_ws_multi_street_and_other_zone_obreb():
    recs = {r["id"]: r for r in baipp_records("baipp-ws-2026-09-ytd-ezd-signs", "2026")}
    assert list(recs) == ["ws-am-6730-2026-2", "ws-am-6730-2025-27"]
    port = recs["ws-am-6730-2025-27"]
    assert (port["procedure"], port["outcome"], port["date"], port["period"]) == ("wz", None, "2026-02-25", "2026-02")
    assert [p["street"] for p in port["address"]["parts"]] == ["Sokola", "Zamoście", "Wybrzeże Szczecińskie"]
    assert port["parcels"] == [{"number": "18/5", "obreb": "4-15-08", "partial": False, "id": "146507_8.1508.18/5"}]
    skalskiego = recs["ws-am-6730-2026-2"]
    assert [p["street"] for p in skalskiego["address"]["parts"]] == ["Skalskiego",   # normalize.STREET_CANONICAL
                                                                     "Nowaka-Jeziorańskiego"]
    assert len(skalskiego["parcels"]) == 7


def test_normalize_row_baipp_road_stretch_and_cut_obreb():
    recs = {r["id"]: r for r in baipp_records("baipp-cp-c2-2022-cut-obreb", "2022")}
    assert sorted(recs) == [f"cp-am-2022-c2-{i}" for i in range(1, 7)]
    tram = recs["cp-am-2022-c2-2"]
    assert tram["outcome"] == "discontinued" and tram["address"]["approximate"] is True
    assert tram["address"]["display"] == ("osiedle Saska Kępa na odcinku al. Stanów Zjednoczonych – gen. Skalskiego / "
                                          "Bora Komorowskiego")
    assert [p["street"] for p in tram["address"]["parts"]] == ["Stanów Zjednoczonych", "Skalskiego",
                                                               "Bora-Komorowskiego"]
    assert "26/10" not in [p["number"] for p in tram["parcels"]]       # listed before the cut obręb "3-05-2"
    assert any("niepełny numer obrębu „3-05-2”" in w and "26/10, 26/11" in w for w in tram["warnings"])
    assert recs["cp-am-2022-c2-6"]["outcome"] == "granted"


def test_normalize_row_baipp_series_from_number_and_note_rows():
    (rec,) = baipp_records("baipp-cp-c1-2021-note-row", "2021")
    assert (rec["id"], rec["outcome"]) == ("cp-am-2021-c1-3", "granted")


def test_normalize_row_date_outside_the_file_year_is_flagged():
    row = {"sheet": "S", "row": 9, "warnings": [], "raw": [],
           "cells": {"lp": 1, "number": "5/PRD/21\nwz", "date": "2020-12-30", "applicant": "x",
                     "investment": "budowa", "address": "Grochowska 1\nPraga Południe", "parcels": "1 3-02-03"}}
    rec = normalize_row(row, "wz", {"file_id": "f", "page_id": "p", "period": "2021", "dataset": "baipp"})
    assert rec["period"] == "2020-12" and any("spoza roku zestawienia (2021)" in w for w in rec["warnings"])
    assert any("rok w numerze decyzji (2021) różni się od roku daty (2020)" in w for w in rec["warnings"])


def test_baipp_ids_never_collide_with_district_ids():
    district = re.compile(r"^(?:wz|cp)-\d{4}-\d+|^ws-\d{4}-\d{4}-\d+|^(?:wz|cp|ws|x)-f")
    records = [r for name, *_ in BAIPP_FIXTURES for r in baipp_records(name, "2026")]
    assert records and all(not district.match(r["id"]) and "-am-" in r["id"] for r in records)


@pytest.mark.parametrize("raw, expected, warning", [
    ("cz. 7/5; 3-05-24", [("7/5", "3-05-24", True)], None),
    ("14 3-02-03 8; 10; 11; 12; 21 3-02-02",
     [("14", "3-02-03", False)] + [(n, "3-02-02", False) for n in ("8", "10", "11", "12", "21")], None),
    ("18/5; 4-15-08", [("18/5", "4-15-08", False)], None),
    ("26/10 i 26/11 3-05-2;\n6/14, 6/1 3-01-20", [("6/14", "3-01-20", False), ("6/1", "3-01-20", False)],
     "niepełny numer obrębu"),
])
def test_parse_parcels_city_wide(raw, expected, warning):
    parcels, warnings = parse_parcels(raw, any_obreb=True)
    assert [(p["number"], p["obreb"], p["partial"]) for p in parcels] == expected
    assert all(p["id"].startswith("146507_8.") for p in parcels)       # corrected later by the ULDK probe
    assert (warning is None and not warnings) or any(warning in w for w in warnings)


def test_parse_parcels_district_mode_ignores_other_zones():
    parcels, warnings = parse_parcels("18/5; 4-15-08")
    assert parcels == [] and warnings


def test_parse_address_road_stretch():
    a = parse_address("al. Stanów Zjednoczonych na odcinku od ul. Bajońskiej do ul. Międzynarodowej")
    assert [p["street"] for p in a["parts"]] == ["Stanów Zjednoczonych", "Bajońska", "Międzynarodowa"]
    assert a["approximate"] is True
    a = parse_address(" al.  Jerzego Waszyngtona \n na odcinku   Międzynarodowa - Kinowa")
    assert a["display"] == "al. Jerzego Waszyngtona na odcinku Międzynarodowa – Kinowa"
    assert [p["street"] for p in a["parts"]] == ["Waszyngtona", "Międzynarodowa", "Kinowa"]


# =============================================================================
# merge_duplicates: snapshots of the same yearly list
# =============================================================================


def test_snapshots_of_the_same_year_merge_into_one_record_oldest_source_first():
    older = baipp_records("baipp-wz-2026-09-ytd-numbering-without-1", "2026", file_id="aaaaaaaaaa")
    newer = baipp_records("baipp-wz-2026-09-ytd-numbering-without-1", "2026", file_id="bbbbbbbbbb")
    (merged,) = merge_duplicates(older + newer)
    assert merged["id"] == "wz-am-2026-95"
    assert [s["file_id"] for s in merged["sources"]] == ["aaaaaaaaaa", "bbbbbbbbbb"]
    assert merged["warnings"] == []


def test_edited_snapshot_wins_with_a_warning():
    (old,) = baipp_records("baipp-wz-2026-09-ytd-numbering-without-1", "2026", file_id="aaaaaaaaaa")
    new = json.loads(json.dumps(old))
    new["sources"][0]["file_id"] = "bbbbbbbbbb"
    new["investment"] = "odmowa ustalenia warunków zabudowy (poprawiony opis)"
    (merged,) = merge_duplicates([old, new])
    assert merged["investment"] == "odmowa ustalenia warunków zabudowy (poprawiony opis)"
    assert [s["file_id"] for s in merged["sources"]] == ["aaaaaaaaaa", "bbbbbbbbbb"]
    assert any("różni się między wersjami zestawienia BAiPP (opis inwestycji)" in w for w in merged["warnings"])
    assert old["investment"].startswith("odmowa ustalenia warunków zabudowy dotyczy")   # inputs not mutated


def test_same_id_twice_in_one_file_is_still_split():
    (rec,) = baipp_records("baipp-wz-2026-09-ytd-numbering-without-1", "2026")
    twin = json.loads(json.dumps(rec))
    twin["sources"][0]["row"] = 200
    twin["date"] = "2026-09-30"
    out = merge_duplicates([rec, twin])
    assert [r["id"] for r in out] == ["wz-am-2026-95", "wz-am-2026-95-2"]


# =============================================================================
# Parcels outside the district (ULDK probe)
# =============================================================================


ULDK_MISS = "-1 brak wyników\n"


def uldk_polygon(pid: str, lon: float = 21.02, lat: float = 52.26, size: float = 0.0005) -> str:
    return (f"0\n{pid}|SRID=4326;POLYGON(({lon} {lat},{lon + size} {lat},{lon + size} {lat + size},"
            f"{lon} {lat + size},{lon} {lat}))\n")


class FakeUldk:
    def __init__(self, known: dict[str, str]):
        self.known = known
        self.asked: list[str] = []

    def get(self, url, params=None, timeout=None):
        assert url == gc.ULDK_URL
        self.asked.append(params["id"])
        return FakeResponse(url, 200, self.known.get(params["id"], ULDK_MISS))


def _record(rid, parcels):
    return {"id": rid, "dataset": "baipp", "warnings": [],
            "parcels": [{"number": n, "obreb": o, "partial": False,
                         "id": f"146507_8.{o[2:].replace('-', '')}.{n}"} for n, o in parcels]}


def _permit(rid, code, number, method=None, organ="Prezydent Miasta Warszawa", decision="1/2020"):
    return {"id": rid, "dataset": "rwdz", "geo": {"method": method}, "number": decision,
            "sources": [{"raw": [["nazwa_organu", organ]]}],
            "warnings": [f"obręb „{code}” nie występuje na Pradze-Południe (nie ma go w ewidencji gruntów dzielnicy) – "
                         "możliwy błąd w rejestrze", "inna uwaga"],
            "parcels": [{"number": number, "obreb": code, "partial": False, "id": f"146507_8.{code}.{number}"}]}


REGION_OK = "0\n{}|\n"


def test_permits_with_a_precinct_of_another_district_are_moved_out(tmp_path):
    from pipeline.normalize import is_pp_obreb
    # the 4-digit codes repeat across zones (Praga-Północ 4-15-07 and Wawer 3-15-07): the case names the unit
    zgn = _permit("pb-st-mz-pg-2016-6120", "1507", "3", organ="Prezydent Miasta Warszawa - Praga Pn.")
    kijowska = _permit("pb-st-mz-wa-2022-11434", "1406", "36/2", decision="56/PRN/2022")   # parcel gone
    kaweczynska = _permit("pb-st-mz-wa-2019-27453", "1406", "11", decision="272/PRN/2019")  # 11 exists in Wawer only
    wawer = _permit("pb-st-mz-pa-2021-49120", "1122", "117/50", organ="Prezydent Miasta Warszawa - Praga Płd.")
    serocka = _permit("pb-st-mz-pa-2018-30453", "1412", "89", method="address")   # located in the district: kept
    twice = _permit("pb-x", "0002", "39/6", organ="Wojewoda Mazowiecki")   # no hint, precinct in two units
    single = _permit("pb-y", "1111", "1", organ="Wojewoda Mazowiecki")     # no hint, precinct in one unit only
    known = {pid: uldk_polygon(pid) for pid in ("146508_8.1507.3", "146514_8.1507.3", "146514_8.1406.11",
                                                "146514_8.1122.117/50", "146508_8.1412.89", "146505_8.0002.39/6",
                                                "146518_8.0002.39/6", "146514_8.1111.1")}
    known.update({rid: REGION_OK.format(rid) for rid in ("146508_8.1507", "146514_8.1507", "146508_8.1406",
                                                         "146514_8.1406", "146514_8.1122", "146505_8.0002",
                                                         "146518_8.0002", "146514_8.1111")})
    session = FakeUldk(known)
    recs = [zgn, kijowska, kaweczynska, wawer, serocka, twice, single]
    out = baipp.check_out_of_district_permits(recs, is_pp_obreb=is_pp_obreb, cache_dir=str(tmp_path), online=True,
                                              session=session)
    assert out == {"pb-st-mz-pg-2016-6120", "pb-st-mz-wa-2022-11434", "pb-st-mz-wa-2019-27453",
                   "pb-st-mz-pa-2021-49120", "pb-y"}
    assert (zgn["parcels"][0]["id"], zgn["parcels"][0]["obreb"]) == ("146508_8.1507.3", "4-15-07")
    assert zgn["warnings"] == ["inna uwaga", "działka poza dzielnicą Praga-Południe według ewidencji gruntów (ULDK): "
                               "3 obr. 1507: Praga-Północ, 146508_8.1507.3 – w rejestrze z identyfikatorem jednostki "
                               "146507_8 (Praga-Południe); nie zaznaczono jej na mapie dzielnicy"]
    assert (kijowska["parcels"][0]["id"], kijowska["parcels"][0]["obreb"]) == ("146508_8.1406.36/2", "4-14-06")
    assert "obręb 146508_8.1406 (Praga-Północ), działki już w nim nie ma" in kijowska["warnings"][-1]
    assert kaweczynska["parcels"][0]["id"] == "146508_8.1406.11"          # not Wawer's parcel 11
    assert (wawer["parcels"][0]["id"], wawer["parcels"][0]["obreb"]) == ("146514_8.1122.117/50", "3-11-22")
    assert single["parcels"][0]["id"] == "146514_8.1111.1"
    assert serocka["parcels"][0]["id"] == "146507_8.1412.89" and len(serocka["warnings"]) == 2
    assert not any(a.startswith("146508_8.1412") for a in session.asked)  # a located record is never probed
    assert twice["parcels"][0]["id"] == "146507_8.0002.39/6" and len(twice["warnings"]) == 2
    # offline, from the cache: the same answer without requests
    again = _permit("pb-st-mz-pg-2016-6120", "1507", "3", organ="Prezydent Miasta Warszawa - Praga Pn.")
    n = len(session.asked)
    assert baipp.check_out_of_district_permits([again], is_pp_obreb=is_pp_obreb, cache_dir=str(tmp_path),
                                               online=False) == {"pb-st-mz-pg-2016-6120"}
    assert len(session.asked) == n and again["parcels"] == zgn["parcels"]
    # nothing cached and offline: nothing moves
    assert baipp.check_out_of_district_permits([_permit("pb-z", "1280", "1")], is_pp_obreb=is_pp_obreb,
                                               cache_dir=str(tmp_path), online=False) == set()


def test_check_out_of_district_probes_the_same_zone_and_flags_the_record(tmp_path):
    port = _record("ws-am-6730-2025-27", [("18/5", "4-15-08")])
    inside = _record("wz-am-2026-95", [("7/5", "3-05-24")])
    gone = _record("wz-am-2023-110", [("3", "3-02-03")])           # subdivided PP parcel: probed, never moved
    district = {"id": "wz-2026-1", "dataset": "district", "warnings": [],
                "parcels": [{"number": "1", "obreb": "3-02-03", "id": "146507_8.0203.1"}]}
    session = FakeUldk({"146508_8.1508.18/5": uldk_polygon("146508_8.1508.18/5"),
                        "146507_8.0524.7/5": uldk_polygon("146507_8.0524.7/5", 21.11, 52.23),
                        "146505_8.0203.3": uldk_polygon("146505_8.0203.3")})   # Mokotów's 1-02-03: other zone
    out = baipp.check_out_of_district([port, inside, gone, district], cache_dir=str(tmp_path), online=True,
                                      session=session)
    assert out == {"ws-am-6730-2025-27"}
    assert port["parcels"][0]["id"] == "146508_8.1508.18/5"
    (warning,) = port["warnings"]
    assert warning.startswith("działka poza dzielnicą Praga-Południe według ewidencji gruntów (ULDK): 18/5 obr. "
                              "4-15-08: Praga-Północ, 146508_8.1508.18/5 – ")
    assert inside["warnings"] == [] and gone["warnings"] == [] and gone["parcels"][0]["id"] == "146507_8.0203.3"
    assert "146505_8.0203.3" not in session.asked                    # never probed outside the obręb zone
    assert {a for a in session.asked if a.endswith(".0203.3")} == {"146507_8.0203.3", "146509_8.0203.3",
                                                                    "146514_8.0203.3"}  # own unit, then its zone
    assert "146515_8.0203.3" not in session.asked                    # Wesoła is zone 8 ("8-06-02")
    assert not any(a.startswith("146507_8.0203.1") for a in session.asked)  # district records are not probed
    cache = gc.load_cache(str(tmp_path / "uldk.json"))
    assert cache["146508_8.1508.18/5"]["status"] == "ok" and cache["146511_8.1508.18/5"]["status"] == "missing"
    assert "146509_8.1508.18/5" not in cache   # Rembertów is zone 3, never asked for a zone-4 obręb

    # offline, from the cache only: the same answer, no request
    again = _record("ws-am-6730-2025-27", [("18/5", "4-15-08")])
    assert baipp.check_out_of_district([again], cache_dir=str(tmp_path), online=False) == {"ws-am-6730-2025-27"}
    assert again["parcels"][0]["id"] == "146508_8.1508.18/5"


def test_check_out_of_district_a_typo_in_a_zone_3_obreb_is_not_matched_in_wesola(tmp_path):
    # wz-am-2023-114: "77 3-06-02" (typo for 3-06-01); Wesoła's precinct 146515_8.0602 is "8-06-02"
    rec = _record("wz-am-2023-114", [("77", "3-06-02")])
    session = FakeUldk({"146515_8.0602.77": uldk_polygon("146515_8.0602.77")})
    assert baipp.check_out_of_district([rec], cache_dir=str(tmp_path), online=True, session=session) == set()
    assert rec["warnings"] == [] and "146515_8.0602.77" not in session.asked


def test_check_out_of_district_offline_without_cache_changes_nothing(tmp_path):
    rec = _record("ws-am-6730-2025-27", [("18/5", "4-15-08")])
    assert baipp.check_out_of_district([rec], cache_dir=str(tmp_path), online=False) == set()
    assert rec["warnings"] == [] and rec["parcels"][0]["id"] == "146507_8.1508.18/5"


# =============================================================================
# Build integration
# =============================================================================


SQUARE = {"type": "Polygon", "coordinates": [[[21.03, 52.20], [21.13, 52.20], [21.13, 52.265], [21.03, 52.265],
                                              [21.03, 52.20]]]}


def _uldk_ok(lon, lat, size=0.0005):
    geometry = {"type": "Polygon", "coordinates": [[[lon, lat], [lon + size, lat], [lon + size, lat + size],
                                                    [lon, lat + size], [lon, lat]]]}
    return {"status": "ok", "geometry": geometry, "centroid": [lon + size / 2, lat + size / 2],
            "bbox": [lon, lat, lon + size, lat + size], "area_m2": 1900.0, "fetched_at": NOW}


def make_baipp_project(tmp_path: Path) -> build.Paths:
    """Two monthly snapshots of the 2026 WZ list (the second with an edited row), the 2026 ws list and the 2020 WZ
    list, with ULDK answers cached (the Port Praski parcel only exists in Praga-Północ)."""
    root = tmp_path / "proj"
    raw = root / "data" / "raw"
    raw.mkdir(parents=True)
    edited = raw / "bbbbbbbbbb.xlsx"
    shutil.copyfile(fx("baipp-wz-2026-09-ytd-numbering-without-1"), raw / "aaaaaaaaaa.xlsx")
    wb = openpyxl.load_workbook(fx("baipp-wz-2026-09-ytd-numbering-without-1"))
    wb.worksheets[0]["H103"] = "odmowa ustalenia warunków zabudowy dla budynku wielorodzinnego (korekta)"
    wb.save(edited)
    shutil.copyfile(fx("baipp-ws-2026-09-ytd-ezd-signs"), raw / "cccccccccc.xlsx")
    shutil.copyfile(fx("baipp-wz-2020-merged-header"), raw / "dddddddddd.xlsx")

    def entry(fid, category, period, page, last_modified, removed=None):
        f = {"file_id": fid, "source": "baipp", "url": f"{SITE}/documents/x/{fid}.xlsx?t=1", "document_path": f"/d/{fid}",
             "filename": f"{fid}.xlsx", "ext": "xlsx", "category": category, "period": period, "page_ids": [page],
             "local_path": f"data/raw/{fid}.xlsx", "sha256": None, "size": None, "downloaded_at": NOW,
             "last_modified": last_modified, "history": [], "series": None}
        if removed:
            f["removed_from_source_at"] = removed
        return f

    files = [entry("aaaaaaaaaa", "wz", "2026", "postepowania-i-decyzje", "2026-09-10T00:00:00Z", removed="2026-10-12"),
             entry("bbbbbbbbbb", "wz", "2026", "postepowania-i-decyzje", "2026-10-12T09:48:10Z"),
             entry("cccccccccc", "ws", "2026", "postepowania-i-decyzje", "2026-10-12T09:48:11Z"),
             entry("dddddddddd", "wz", "2020", "rejestry2020", "2021-10-17T19:12:11Z")]
    pages = [
        {"page_id": "postepowania-i-decyzje", "url": BAIPP_URL, "title": "Informacja o wszczętych postępowaniach",
         "period": "2026", "published_at": None, "summary": None, "index_page": None, "source": "baipp", "warnings": [],
         "blocks": [{"text": "1. Informacja o wszczętych …", "category": "ws", "declared_count": None,
                     "file_ids": ["cccccccccc"], "series": None},
                    {"text": "2. Informacja o wydanych decyzjach w sprawie warunków zabudowy X.2026", "category": "wz",
                     "declared_count": None, "file_ids": ["bbbbbbbbbb"], "series": None}]},
        {"page_id": "rejestry2020", "url": SITE + "/-/rejestry2020", "title": "… (sierpień - grudzień 2020 r.)",
         "period": "2020", "published_at": "2022-02-16T14:09:00Z", "summary": None, "index_page": None,
         "source": "baipp", "warnings": [],
         "blocks": [{"text": "2. … warunków zabudowy - VIII-XII 2020", "category": "wz", "declared_count": None,
                     "file_ids": ["dddddddddd"], "series": None}]},
    ]
    crawl_mod.write_manifest({"generated_at": NOW, "index_url": "https://src.example/index", "pages": pages,
                              "files": files}, root / "data" / "manifest.json")
    cache = root / "data" / "cache"
    cache.mkdir(parents=True)
    uldk = {"146507_8.0524.7/5": _uldk_ok(21.11, 52.23), "146507_8.1508.18/5": {"status": "missing", "fetched_at": NOW},
            "146508_8.1508.18/5": _uldk_ok(21.02, 52.27), "146511_8.1508.18/5": {"status": "missing", "fetched_at": NOW},
            "146503_8.1508.18/5": {"status": "missing", "fetched_at": NOW}}
    (cache / "uldk.json").write_text(json.dumps(uldk))
    (cache / "uug.json").write_text("{}")
    site_data = root / "site" / "data"
    site_data.mkdir(parents=True)
    (site_data / "areas.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {"slug": "_district", "label": "Praga-Południe"}, "geometry": SQUARE},
        {"type": "Feature", "properties": {"slug": "goclaw", "label": "Gocław"}, "geometry": SQUARE}]}))
    return build.Paths(root)


def test_build_with_baipp_files(tmp_path):
    paths = make_baipp_project(tmp_path)
    result = build.run_build(paths, crawl=False, online=False, site_url="https://ex.org/", now=lambda: NOW)
    assert result.exit_code == build.EXIT_OK
    records_doc = json.loads(paths.records_json.read_text(encoding="utf-8"))
    sources_doc = json.loads(paths.sources_json.read_text(encoding="utf-8"))
    by_id = {r["id"]: r for r in records_doc["records"]}
    assert sorted(by_id) == sorted(["wz-am-2026-95", "ws-am-6730-2026-2", "ws-am-6730-2025-27", "wz-am-2020-98",
                                    "wz-am-2020-110", "wz-am-2020-123", "wz-am-2020-154"])
    assert records_doc["counts_by_kind"] == {"wz": 5, "cp": 0, "ws": 2}
    assert records_doc["datasets"]["baipp"]["url"] == BAIPP_URL and "city" in records_doc["issuers"]

    # two snapshots of one yearly list → one record, both rows as sources (oldest first), newest values
    wz95 = by_id["wz-am-2026-95"]
    assert [(s["file_id"], s["row"]) for s in wz95["sources"]] == [("aaaaaaaaaa", 103), ("bbbbbbbbbb", 103)]
    assert wz95["investment"].endswith("(korekta)")
    assert any("różni się między wersjami zestawienia BAiPP" in w for w in wz95["warnings"])
    assert (wz95["dataset"], wz95["issuer"], wz95["outcome"], wz95["period"]) == ("baipp", "city", "refusal", "2026-09")
    assert wz95["geo"]["method"] == "parcel" and wz95["area"] == "goclaw"

    # labelled Praga Południe, parcel in Praga-Północ: kept, flagged, not on the map
    port = by_id["ws-am-6730-2025-27"]
    assert port["geo"]["method"] is None and port["geo"]["lat"] is None and port["area"] is None
    assert "poza dzielnicą" in port["geo"]["note"] and port["parcels"][0]["id"] == "146508_8.1508.18/5"
    assert any(w.startswith(baipp.OUTSIDE_WARNING) for w in port["warnings"])
    parcels = json.loads(paths.parcels_geojson.read_text(encoding="utf-8"))
    assert "146508_8.1508.18/5" not in {f["id"] for f in parcels["features"]}

    # sources.json: city-wide rows vs Praga-Południe rows, snapshot kept although no longer linked
    files = {f["file_id"]: f for f in sources_doc["files"]}
    assert (files["aaaaaaaaaa"]["rows_total"], files["aaaaaaaaaa"]["parsed_count"]) == (101, 1)
    assert (files["cccccccccc"]["rows_total"], files["cccccccccc"]["parsed_count"]) == (85, 2)
    assert (files["dddddddddd"]["rows_total"], files["dddddddddd"]["parsed_count"]) == (79, 4)
    assert files["aaaaaaaaaa"]["record_ids"] == ["wz-am-2026-95"] and files["aaaaaaaaaa"]["layout"] == "baipp-wz"
    assert files["aaaaaaaaaa"]["archive_path"] == "sources/aaaaaaaaaa.xlsx"
    blocks = {b["file_ids"][0]: b for p in sources_doc["pages"] for b in p["blocks"]}
    assert blocks["bbbbbbbbbb"]["parsed_count"] == 1 and blocks["bbbbbbbbbb"]["count_mismatch"] is False
    assert sum(len(r["sources"]) for r in records_doc["records"]) == sum(f["parsed_count"] for f in files.values())

    report = paths.report.read_text(encoding="utf-8")
    assert "## BAiPP (city architecture bureau) lists" in report
    assert "| 101 | 1 | 1 |" in report and "no longer linked since 2026-10-12" in report
    assert "[`ws-am-6730-2025-27`](https://ex.org/#/wpis/ws-am-6730-2025-27)" in report

    # deterministic: a rebuild changes nothing
    again = build.run_build(paths, crawl=False, online=False, site_url="https://ex.org/", now=lambda: "2099-01-01T00:00:00Z")
    assert again.written == [] and again.generated_at == NOW
